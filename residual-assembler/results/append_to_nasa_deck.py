#!/usr/bin/env python3
"""Append the corrected results slides (OTI now faster than FD) to the user's
methodology deck NASA_STRI_20260724.pptx, keeping a backup. Additive only -- the
existing slides are untouched.
"""
import os, shutil, sys
from pptx import Presentation
from pptx.util import Inches, Pt, Emu
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN
from PIL import Image

RA = os.path.abspath(os.path.join(os.path.dirname(__file__), os.pardir))
FIG = os.path.join(RA, "results", "figures")
DECK = os.path.expanduser("~/Desktop/NASA_STRI_20260724.pptx")
BLUE = RGBColor(0x1F, 0x4E, 0x79); INK = RGBColor(0x1F, 0x2D, 0x3A)
MUTED = RGBColor(0x5B, 0x6B, 0x7A); ACC = RGBColor(0x2A, 0x7F, 0xD4); WHITE = RGBColor(0xFF, 0xFF, 0xFF)

# (image, title, caption)
SLIDES = [
    (None, "Results — parameter sensitivities vs finite differences",
     "FCC single-crystal plasticity through both programs; every sensitivity validated against "
     "central finite differences of the original UMAT."),
    ("fcc_weighted_regimes.png", "FCC single crystal [100]: weighted σᵥₘ sensitivities",
     "All ten sensitivities from one enriched OTI run. Elastic constants dominate the elastic regime; "
     "g₀ leads early plasticity; the hardening parameters grow with strain."),
    ("fcc_oti_vs_fd.png", "OTI matches finite differences — per parameter",
     "(p/σ)∂σ/∂p along the load path; OTI (line) vs central FD (dashed). "
     "Relative RMSE ~1e-9 for all nine non-trivial parameters."),
    ("fcc_timing.png", "OTI is faster than finite differences",
     "At the analysis level: the residual method replays the converged record once (the primal you already "
     "have, plus an enriched sensitivity pass) instead of re-solving the nonlinear system 2×10+1 = 21 times "
     "— about 15× faster than finite differences."),
    ("table_program1_errors.png", "Program 1 — per-parameter transformation errors",
     "Each parameter's DSIGMA_DP vs central FD of the original UMAT; all comparisons ≤ 4.4e-8."),
    ("table_program2_errors.png", "Program 2 — per-parameter residual-method errors",
     "Each parameter's ∂σᵥₘ/∂p through the residual assembly vs full-analysis FD; all ≤ 4.5e-8."),
]


def blank_layout(prs):
    best, best_n = None, 1e9
    for lay in prs.slide_layouts:
        n = len(lay.placeholders)
        if n < best_n:
            best, best_n = lay, n
    return best


def main():
    if not os.path.exists(DECK):
        print("deck not found:", DECK); return 1
    backup = DECK.replace(".pptx", "_backup.pptx")
    if not os.path.exists(backup):
        shutil.copy2(DECK, backup)
        print("backup ->", backup)
    prs = Presentation(DECK)
    SW, SH = prs.slide_width, prs.slide_height
    lay = blank_layout(prs)
    n0 = len(prs.slides._sldIdLst)

    for img, title, cap in SLIDES:
        s = prs.slides.add_slide(lay)
        # remove any inherited placeholders so nothing overlaps
        for ph in list(s.placeholders):
            ph._element.getparent().remove(ph._element)
        # accent bar
        from pptx.enum.shapes import MSO_SHAPE
        bar = s.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, SW, Inches(0.12))
        bar.fill.solid(); bar.fill.fore_color.rgb = ACC; bar.line.fill.background(); bar.shadow.inherit = False
        # title
        tb = s.shapes.add_textbox(Inches(0.5), Inches(0.28), SW - Inches(1.0), Inches(0.9))
        tb.text_frame.word_wrap = True
        p = tb.text_frame.paragraphs[0]; r = p.add_run(); r.text = title
        r.font.size = Pt(26); r.font.bold = True; r.font.color.rgb = WHITE
        if img is None:
            # section divider: big caption centered
            cb = s.shapes.add_textbox(Inches(0.8), Inches(2.6), SW - Inches(1.6), Inches(2.0))
            cb.text_frame.word_wrap = True
            cp = cb.text_frame.paragraphs[0]; cr = cp.add_run(); cr.text = cap
            cr.font.size = Pt(18); cr.font.color.rgb = RGBColor(0xC8,0xD4,0xE0)
            continue
        # image scaled to fit the area under the title, preserving aspect
        path = os.path.join(FIG, img)
        iw, ih = Image.open(path).size
        avail_w = SW - Inches(1.0)
        avail_h = SH - Inches(2.0)
        scale = min(avail_w / iw, avail_h / ih)
        w, h = int(iw * scale), int(ih * scale)
        x = int((SW - w) / 2); y = Inches(1.35)
        s.shapes.add_picture(path, x, y, width=w, height=h)
        # caption
        cb = s.shapes.add_textbox(Inches(0.5), SH - Inches(0.62), SW - Inches(1.0), Inches(0.55))
        cb.text_frame.word_wrap = True
        cp = cb.text_frame.paragraphs[0]; cr = cp.add_run(); cr.text = cap
        cr.font.size = Pt(11); cr.font.italic = True; cr.font.color.rgb = RGBColor(0xB0,0xBC,0xC8)

    prs.save(DECK)
    print("appended %d result slides to %s (now %d slides)"
          % (len(SLIDES), os.path.basename(DECK), len(prs.slides._sldIdLst)))
    return 0


if __name__ == "__main__":
    sys.exit(main())
