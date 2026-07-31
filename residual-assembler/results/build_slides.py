#!/usr/bin/env python3
"""Build the two-program results slide deck (.pptx) in the NASA STRI style.

Covers Program 1 (UMAT->OTI transformation) and Program 2 (residual assembler):
methodology, the required validations, and every plot. All numbers are the actual
validated results (see results/figures/*.json and the material contracts).

    python results/build_slides.py   ->   results/NASA_STRI_residual_assembler.pptx
"""
import os, sys
from pptx import Presentation
from pptx.util import Inches, Pt, Emu
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.shapes import MSO_SHAPE

RA = os.path.abspath(os.path.join(os.path.dirname(__file__), os.pardir))
FIG = os.path.join(RA, "results", "figures")

INK   = RGBColor(0x1F, 0x2D, 0x3A)
BLUE  = RGBColor(0x1F, 0x4E, 0x79)
ACCENT= RGBColor(0x2A, 0x7F, 0xD4)
MUTED = RGBColor(0x5B, 0x6B, 0x7A)
GREENB= RGBColor(0xE2, 0xEF, 0xDA)
GREENT= RGBColor(0x37, 0x56, 0x23)
CODEB = RGBColor(0xF2, 0xF3, 0xF5)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)
LINE  = RGBColor(0xD9, 0xDE, 0xE3)
HDRBG = RGBColor(0x4A, 0x6D, 0x8C)

EMU_W, EMU_H = Inches(13.333), Inches(7.5)


def _txt(tf, runs, size, color=INK, bold=False, align=PP_ALIGN.LEFT, space_after=6, italic=False):
    """runs: list of (text, {overrides}) or a plain string -> one paragraph per entry."""
    first = True
    for item in runs:
        p = tf.paragraphs[0] if first else tf.add_paragraph()
        first = False
        p.alignment = align; p.space_after = Pt(space_after)
        segs = item if isinstance(item, list) else [item]
        for seg in segs:
            if isinstance(seg, tuple):
                text, ov = seg
            else:
                text, ov = seg, {}
            r = p.add_run(); r.text = text
            r.font.size = Pt(ov.get("size", size)); r.font.bold = ov.get("bold", bold)
            r.font.italic = ov.get("italic", italic)
            r.font.color.rgb = ov.get("color", color); r.font.name = ov.get("font", "Calibri")
        if isinstance(item, list) is False and isinstance(item, str):
            pass
    return tf


def _box(slide, x, y, w, h, fill=None, line=None, line_w=0.75):
    sp = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, x, y, w, h)
    sp.shadow.inherit = False
    if fill is None:
        sp.fill.background()
    else:
        sp.fill.solid(); sp.fill.fore_color.rgb = fill
    if line is None:
        sp.line.fill.background()
    else:
        sp.line.color.rgb = line; sp.line.width = Pt(line_w)
    return sp


def _rrect(slide, x, y, w, h, fill, line=None):
    sp = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, x, y, w, h)
    sp.shadow.inherit = False
    sp.fill.solid(); sp.fill.fore_color.rgb = fill
    if line is None:
        sp.line.fill.background()
    else:
        sp.line.color.rgb = line; sp.line.width = Pt(1)
    return sp


def _blank(prs):
    s = prs.slides.add_slide(prs.slide_layouts[6])
    return s


def _header(slide, kicker, title):
    _box(slide, 0, 0, EMU_W, Inches(0.14), fill=ACCENT)          # top accent bar
    if kicker:
        tb = slide.shapes.add_textbox(Inches(0.55), Inches(0.34), Inches(12.2), Inches(0.3))
        _txt(tb.text_frame, [[(kicker.upper(), {"color": ACCENT, "bold": True, "size": 12})]], 12, space_after=0)
    tb = slide.shapes.add_textbox(Inches(0.55), Inches(0.6), Inches(12.2), Inches(0.9))
    _txt(tb.text_frame, [[(title, {"color": BLUE, "bold": True, "size": 30})]], 30, space_after=0)


def _bullets(slide, x, y, w, h, items, size=16):
    tb = slide.shapes.add_textbox(x, y, w, h); tf = tb.text_frame; tf.word_wrap = True
    first = True
    for it in items:
        p = tf.paragraphs[0] if first else tf.add_paragraph(); first = False
        p.space_after = Pt(9); p.space_before = Pt(0)
        lvl = it[2] if len(it) > 2 else 0
        p.level = lvl
        bullet = "•  " if lvl == 0 else "–  "
        r = p.add_run(); r.text = bullet; r.font.size = Pt(size); r.font.color.rgb = ACCENT if lvl == 0 else MUTED; r.font.bold = True
        for seg in (it[0] if isinstance(it[0], list) else [it[0]]):
            text, ov = seg if isinstance(seg, tuple) else (seg, {})
            r = p.add_run(); r.text = text
            r.font.size = Pt(ov.get("size", size)); r.font.bold = ov.get("bold", False)
            r.font.color.rgb = ov.get("color", INK); r.font.name = ov.get("font", "Calibri")
            r.font.italic = ov.get("italic", False)
    return tb


def _takeaway(slide, x, y, w, text, h=Inches(0.9)):
    _rrect(slide, x, y, w, h, fill=GREENB)
    tb = slide.shapes.add_textbox(x + Inches(0.18), y + Inches(0.1), w - Inches(0.36), h - Inches(0.2))
    tb.text_frame.word_wrap = True
    tb.text_frame.vertical_anchor = MSO_ANCHOR.MIDDLE
    _txt(tb.text_frame, [text if isinstance(text, list) else [(text, {"color": GREENT, "bold": True, "size": 15})]],
         15, color=GREENT, bold=True, space_after=0)


def _figure(slide, path, x, y, w=None, h=None):
    if w and not h:
        return slide.shapes.add_picture(path, x, y, width=w)
    if h and not w:
        return slide.shapes.add_picture(path, x, y, height=h)
    return slide.shapes.add_picture(path, x, y, width=w, height=h)


def _caption(slide, x, y, w, text):
    tb = slide.shapes.add_textbox(x, y, w, Inches(0.5)); tb.text_frame.word_wrap = True
    _txt(tb.text_frame, [[(text, {"color": MUTED, "size": 11, "italic": True})]], 11, space_after=0)


def _codebox(slide, x, y, w, h, lines):
    _rrect(slide, x, y, w, h, fill=CODEB, line=LINE)
    tb = slide.shapes.add_textbox(x + Inches(0.15), y + Inches(0.1), w - Inches(0.3), h - Inches(0.2))
    tf = tb.text_frame; tf.word_wrap = True
    first = True
    for ln, col in lines:
        p = tf.paragraphs[0] if first else tf.add_paragraph(); first = False
        p.space_after = Pt(1)
        r = p.add_run(); r.text = ln; r.font.size = Pt(11.5); r.font.name = "Consolas"
        r.font.color.rgb = col
    return tb


def _table(slide, x, y, w, headers, rows, colw=None, fs=12, rowh=0.32):
    nr, nc = len(rows) + 1, len(headers)
    gtbl = slide.shapes.add_table(nr, nc, x, y, w, Inches(rowh * nr)).table
    if colw:
        total = sum(colw)
        for j, cw in enumerate(colw):
            gtbl.columns[j].width = Emu(int(w * cw / total))
    for j, htext in enumerate(headers):
        c = gtbl.cell(0, j); c.fill.solid(); c.fill.fore_color.rgb = HDRBG
        c.margin_top = Pt(2); c.margin_bottom = Pt(2)
        tf = c.text_frame; tf.word_wrap = True
        for li, line in enumerate(htext.split("\n")):
            p = tf.paragraphs[0] if li == 0 else tf.add_paragraph()
            r = p.add_run(); r.text = line; r.font.size = Pt(fs); r.font.bold = True; r.font.color.rgb = WHITE
    for i, row in enumerate(rows, 1):
        for j, val in enumerate(row):
            c = gtbl.cell(i, j)
            c.fill.solid(); c.fill.fore_color.rgb = WHITE if i % 2 else RGBColor(0xF4, 0xF6, 0xF8)
            c.margin_top = Pt(1); c.margin_bottom = Pt(1)
            tf = c.text_frame; p = tf.paragraphs[0]
            txt, ov = val if isinstance(val, tuple) else (val, {})
            r = p.add_run(); r.text = txt; r.font.size = Pt(ov.get("size", fs))
            r.font.color.rgb = ov.get("color", INK); r.font.bold = ov.get("bold", False)
            r.font.name = ov.get("font", "Calibri")
    return gtbl


def _flow(slide, y, steps):
    """Horizontal workflow strip of pill boxes with arrows."""
    n = len(steps); m = Inches(0.55); gap = Inches(0.18)
    total = EMU_W - 2 * m
    bw = int((total - gap * (n - 1)) / n)
    x = m
    for i, (label, strong) in enumerate(steps):
        _rrect(slide, x, y, bw, Inches(0.62), fill=WHITE if not strong else RGBColor(0xEA, 0xF2, 0xFB), line=LINE)
        tb = slide.shapes.add_textbox(x + Inches(0.05), y + Inches(0.06), bw - Inches(0.1), Inches(0.5))
        tb.text_frame.word_wrap = True; tb.text_frame.vertical_anchor = MSO_ANCHOR.MIDDLE
        _txt(tb.text_frame, [[(ln, {"color": BLUE if strong else INK, "bold": strong, "size": 11})]
                             for ln in label.split("\n")], 11, align=PP_ALIGN.CENTER, space_after=0)
        if i < n - 1:
            ar = slide.shapes.add_textbox(x + bw - Inches(0.02), y + Inches(0.1), gap + Inches(0.04), Inches(0.4))
            _txt(ar.text_frame, [[(u"→", {"color": ACCENT, "bold": True, "size": 16})]], 16, align=PP_ALIGN.CENTER, space_after=0)
        x += bw + gap


# ------------------------------------------------------------------ deck ----
def build():
    prs = Presentation(); prs.slide_width = EMU_W; prs.slide_height = EMU_H

    # 1. TITLE ---------------------------------------------------------------
    s = _blank(prs)
    _box(s, 0, 0, EMU_W, EMU_H, fill=RGBColor(0x14, 0x2A, 0x42))
    _box(s, 0, Inches(4.55), EMU_W, Inches(0.06), fill=ACCENT)
    tb = s.shapes.add_textbox(Inches(0.8), Inches(1.35), Inches(11.9), Inches(2.9))
    tb.text_frame.word_wrap = True
    _txt(tb.text_frame, [
        [("Model-agnostic material sensitivities via OTI", {"color": WHITE, "bold": True, "size": 38})],
        [("A two-program residual-assembler framework, verified end to end", {"color": RGBColor(0x9E, 0xC2, 0xE8), "size": 20})],
    ], 38, space_after=12)
    tb = s.shapes.add_textbox(Inches(0.8), Inches(4.8), Inches(11.9), Inches(1.9))
    tb.text_frame.word_wrap = True
    _txt(tb.text_frame, [
        [("Program 1 — UMAT to OTI source transformation      Program 2 — offline residual assembler", {"color": WHITE, "size": 16, "bold": True})],
        [("Reference throughout: central finite differences of the original UMAT", {"color": RGBColor(0x9E, 0xC2, 0xE8), "size": 14})],
        [("NASA STRI project  ·  Dept. of Mechanical, Aerospace & Industrial Engineering, UT San Antonio", {"color": RGBColor(0x7F, 0xA6, 0xCC), "size": 13})],
    ], 16, space_after=8)

    # 2. WORKFLOW OVERVIEW ---------------------------------------------------
    s = _blank(prs); _header(s, "Overview", "Two programs, one protected boundary")
    _bullets(s, Inches(0.55), Inches(1.5), Inches(12.2), Inches(2.0), [
        [[("JHU develops and maintains the private material-model source; collaborators run production analyses with the compiled binaries.", {})]],
        [[("The residual method needs stress derivatives w.r.t. material parameters — a normal UMAT cannot return them.", {})]],
        [[("Two separately-distributable programs solve this without exposing the source or rerunning the production analysis.", {})]],
    ])
    _flow(s, Inches(3.7), [("JHU private\nUMAT source", True), ("UMAT→OTI\ntransform", False),
                           ("umat_<name>_oti.obj\n+ contract", True), ("collaborator\nFE record", False),
                           ("residual\nassembler", False), ("∂q/∂p\nsensitivities", True)])
    _takeaway(s, Inches(0.55), Inches(4.9), Inches(12.2),
              "Only versioned public contracts and the OTI-enabled .obj cross the boundary — the material source stays private and the production analysis is unchanged.",
              h=Inches(0.95))
    _caption(s, Inches(0.55), Inches(6.1), Inches(12.2),
             "Program 1 is owned by JHU / material developers; Program 2 by the collaborators. The two repos are independent and share only versioned contracts.")

    # 2b. WHAT IS OTI -------------------------------------------------------
    s = _blank(prs); _header(s, "Background", "What is OTI, and why not finite differences?")
    _bullets(s, Inches(0.55), Inches(1.55), Inches(7.2), Inches(4.4), [
        [[("OTI", {"bold": True, "color": BLUE}), (" = Order-Truncated Imaginary numbers — a hypercomplex algebra for exact automatic differentiation (OTILib / HYPAD).", {})]],
        [[("Each chosen material parameter is seeded on its own imaginary ", {}), ("direction", {"italic": True}), ("; every arithmetic operation in the UMAT carries the derivative coefficient automatically.", {})]],
        [[("The derivative flows through everything", {"bold": True, "color": BLUE}), (" — the elastic predictor, the return-mapping Newton iterations, the tangent — with no re-derivation by hand.", {})]],
        [[("∂σ/∂p is read out from the imaginary coefficient — ", {}), ("exact to round-off", {"bold": True}), (", with no truncation or step-size error.", {})]],
        [[("One enriched run seeds many directions at once, returning all requested derivatives together.", {})]],
    ], size=15)
    _codebox(s, Inches(8.0), Inches(1.55), Inches(4.8), Inches(3.2), [
        ("Finite differences        OTI", ACCENT),
        ("---------------------------------------", MUTED),
        ("2N+1 full runs        vs  1 enriched run", INK),
        ("truncation +               round-off", INK),
        ("  step-size error          only", INK),
        ("tune h per param      vs  no tuning", INK),
        ("", INK),
        ("N = # parameters (here N = 6)", MUTED),
    ])
    _takeaway(s, Inches(8.0), Inches(5.0), Inches(4.8),
              "OTI returns the same derivative finite differences approximate — but exactly, and in one run.", h=Inches(1.1))

    # 3. PROGRAM 1 — WHAT IT DOES -------------------------------------------
    s = _blank(prs); _header(s, "Program 1 · UMAT → OTI transformation", "Overload a UMAT to return its derivatives")
    _bullets(s, Inches(0.55), Inches(1.55), Inches(7.0), Inches(4.5), [
        [[("Inputs", {"bold": True, "color": BLUE}), (": the original UMAT source + a compact transformation contract (JSON).", {})]],
        [[("Automatic transform", {"bold": True, "color": BLUE}), (": trace PROPS→STRESS dependencies, promote the required variables/operations to OTI, generate parameter seeding + coefficient extraction.", {})]],
        [[("Output", {"bold": True, "color": BLUE}), (": ", {}), ("umat_<name>_oti.obj", {"font": "Consolas", "size": 14}), (" — one relocatable object with three symbols:", {})]],
        [[("UMAT", {"font": "Consolas", "size": 14}), (" (original, unchanged)", {"color": MUTED})], 0, 1],
        [[("UMAT_OTI_INTERNAL", {"font": "Consolas", "size": 14}), (" (the OTI-promoted routine)", {"color": MUTED})], 0, 1],
        [[("UMAT_OTI_EVAL", {"font": "Consolas", "size": 14}), (" — offline entry returning the derivatives", {"color": MUTED})], 0, 1],
        [[("New outputs", {"bold": True, "color": BLUE}), (": ", {}), ("DSIGMA_DP", {"font": "Consolas", "size": 14}), (" and, for history-dependent models, ", {}), ("DSTATEV_DP", {"font": "Consolas", "size": 14}), (".", {})]],
    ])
    _codebox(s, Inches(7.8), Inches(1.55), Inches(4.95), Inches(3.9), [
        ("SUBROUTINE UMAT_OTI_EVAL(", ACCENT),
        ("  STRESS, STATEV, DDSDDE, STRAN, DSTRAN,", INK),
        ("  TIME,DTIME,TEMP,DTEMP, PROPS,NPROPS,", INK),
        ("  NTENS,NSTATV,NPARAM,", INK),
        ("  DSIGMA_DP,   ! (NTENS , NPARAM)", GREENT),
        ("  DSTATEV_DP,  ! (NSTATV, NPARAM)", GREENT),
        ("  DSIGMA_DP_IN, DSTATEV_DP_IN) ! carry-in", MUTED),
        ("", INK),
        ("* derivatives returned SEPARATELY,", INK),
        ("  never stored in STATEV", INK),
        ("* STATEV_OUT = physical state only", MUTED),
    ])
    _takeaway(s, Inches(0.55), Inches(6.0), Inches(12.2),
              "The standard UMAT interface is preserved; derivative outputs are added. DSIGMA_DP / DSTATEV_DP are dedicated arguments — never packed into STATEV.")

    # 4. PROGRAM 1 — THE COMPACT CONTRACT -----------------------------------
    s = _blank(prs); _header(s, "Program 1 · input", "The transformation contract is compact")
    _bullets(s, Inches(0.55), Inches(1.55), Inches(5.1), Inches(4.5), [
        [[("The developer identifies ", {}), ("only", {"bold": True}),
          (" the UMAT, the parameters to differentiate (name → PROPS index), and the interface sizes.", {})]],
        [[("Parameter ", {}), ("order", {"bold": True}),
          (" fixes the OTI-direction order; ", {}), ("nparam is inferred", {"bold": True}),
          (" as the number of parameters.", {})]],
        [[("A ", {}), ("separate verification case", {"bold": True, "color": BLUE}),
          (" supplies one valid baseline property vector (and optional loading path) — ",
           {}), ("values are where to test, not how to transform.", {})]],
        [[("The program completes the ", {}),
          ("interface, derivative, history, validation, and platform settings automatically.", {"bold": True})]],
    ], size=15)
    _codebox(s, Inches(5.9), Inches(1.5), Inches(6.9), Inches(3.5), [
        ('{  "material": "j2_linear_hardening",', INK),
        ('   "source": "umat_j2.for",', INK),
        ('   "kinematics": "small_strain",', INK),
        ('   "dimensions": {"stress_components": 6,', INK),
        ('                  "material_properties": 4,', INK),
        ('                  "state_variables": 1},', INK),
        ('   "parameters": {"E":1, "nu":2,', GREENT),
        ('                  "SIGY0":3, "H":4},', GREENT),
        ('   "derivatives": {"stress": true,', ACCENT),
        ('                   "statev": "auto"} }', ACCENT),
    ])
    _codebox(s, Inches(5.9), Inches(5.15), Inches(6.9), Inches(1.05), [
        ('verification.json   (separate — the baseline)', MUTED),
        ('{ "props": [210000, 0.30, 250, 2000] }', GREENT),
    ])
    _caption(s, Inches(5.9), Inches(6.35), Inches(6.9),
             "The developer identifies the UMAT, the material parameters to differentiate, and one valid baseline property vector. The program completes the interface, derivative, history, validation, and platform settings automatically.")

    # 5. PROGRAM 1 — VERIFICATION METHODOLOGY -------------------------------
    s = _blank(prs); _header(s, "Program 1 · verification", "Non-circular verification against the original UMAT")
    _bullets(s, Inches(0.55), Inches(1.55), Inches(12.3), Inches(2.7), [
        [[("Primal parity", {"bold": True, "color": BLUE}), (": OTI STRESS, STATEV and consistent tangent DDSDDE are compared against a ", {}), ("separately compiled original UMAT", {"bold": True}), (" — not the transformed binary.", {})]],
        [[("Derivative check", {"bold": True, "color": BLUE}), (": DSIGMA_DP is compared with ", {}), ("central finite differences of the original UMAT", {"bold": True}), (", perturbing PROPS.", {})]],
        [[("Consistent tangent", {"bold": True, "color": BLUE}), (": the OTI DDSDDE is checked against central FD of the original stress w.r.t. DSTRAN (the true algorithmic tangent).", {})]],
        [[("Path dependence", {"bold": True, "color": BLUE}), (": STRESS, STATEV, DSIGMA_DP and DSTATEV_DP are validated over a multi-increment path — elastic → plastic → unload → reload — with the state sensitivity chained across increments.", {})]],
    ], size=15)
    _flow(s, Inches(4.6), [("original UMAT\n(.obj)", False), ("OTI UMAT\n(.obj)", False),
                           ("march same\nstrain path", True), ("compare STRESS,\nSTATEV, DDSDDE", False),
                           ("FD of original\nvs DSIGMA_DP", True)])
    _takeaway(s, Inches(0.55), Inches(5.6), Inches(12.2),
              "The reference is always an independent original-UMAT run — never a self-comparison. The framework passes only when it preserves the response AND produces accurate parameter derivatives.")
    _caption(s, Inches(0.55), Inches(6.65), Inches(12.2),
             "Material-point pass thresholds: stress / STATEV parity ≤ 1e-8; DDSDDE, DSIGMA_DP and DSTATEV_DP vs central FD ≤ 1e-4.")

    # 6. PROGRAM 1 — MATERIAL-POINT VALIDATION TABLE ------------------------
    s = _blank(prs); _header(s, "Program 1 · results", "Material-point validation — all models pass")
    _table(s, Inches(0.4), Inches(1.6), Inches(12.55),
           ["Model", "Constitutive law", "Diff. parameters", "#p", "path", "stress\nparity", "DDSDDE\nvs FD", "DSIGMA_DP\nvs FD", "DSTATEV_DP\nvs FD"],
           [["M1", "Isotropic linear elastic", "E, nu", "2", "–", "2.1e-16", ("5.1e-17", {"color": GREENT, "bold": True}), ("2.3e-10", {"color": GREENT, "bold": True}), "–"],
            ["M2", "Cubic anisotropic elastic", "C11, C12, C44", "3", "–", "0", ("0", {"color": GREENT, "bold": True}), ("2.4e-13", {"color": GREENT, "bold": True}), "–"],
            ["M3", "J2 plasticity, lin. hardening", "E, nu, sig_y0, H", "4", "yes", "8.7e-17", ("2.4e-08", {"color": GREENT, "bold": True}), ("2.2e-08", {"color": GREENT, "bold": True}), ("9.8e-09", {"color": GREENT, "bold": True})],
            ["M5", "Crystal-plasticity flow (Kocks)", "tau0,dG,p,q,gam0,H", "6", "yes", "1.4e-16", ("8.2e-09", {"color": GREENT, "bold": True}), ("2.4e-07", {"color": GREENT, "bold": True}), ("1.6e-07", {"color": GREENT, "bold": True})]],
           colw=[0.72, 1.9, 1.68, 0.4, 0.5, 1.05, 1.15, 1.3, 1.3], fs=11.5, rowh=0.62)
    _takeaway(s, Inches(0.55), Inches(4.7), Inches(12.2),
              [("Stress and state parity is machine precision; every parameter sensitivity matches central finite differences of the original UMAT to a few ×10⁻⁷ or better — including the path-dependent state sensitivity DSTATEV_DP propagated across increments.", {"color": GREENT, "bold": True, "size": 15})],
              h=Inches(1.0))
    _caption(s, Inches(0.55), Inches(5.95), Inches(12.2),
             "M3 and M5 are validated over a 5-increment elastic→plastic→unload→reload path. M5 uses 19 OTI directions (6 DSTRAN + 6 params + 1 state + 6 stress) to chain the state-transition Jacobians.")

    # 7. PROGRAM 1 — 18-UMAT SUITE ------------------------------------------
    s = _blank(prs); _header(s, "Program 1 · results", "Source-transformation test suite (18 / 19 UMATs)")
    _figure(s, os.path.join(FIG, "umat_transform_table.png"), Inches(0.4), Inches(1.5), w=Inches(8.7))
    _bullets(s, Inches(9.3), Inches(1.7), Inches(3.6), Inches(3.5), [
        [[("18 of 19 benchmark UMATs pass", {"bold": True, "color": BLUE})]],
        [[("Different programming approaches and constitutive models: elasticity, plasticity, viscoplasticity, finite-strain.", {})]],
        [[("OTI STRESS/STATEV/DDSDDE compared with the original UMAT in Abaqus.", {})]],
        [[("The lone non-pass (UMAT_VPDCL_R) is an original job that did not converge.", {"color": MUTED})]],
    ], size=13)
    _takeaway(s, Inches(9.3), Inches(4.85), Inches(3.55),
              [("The tool even caught a human error: spin_elas_def's large DDSDDE discrepancy (740) prompted an investigation that found a real bug in the original UMAT.", {"color": GREENT, "bold": True, "size": 13})],
              h=Inches(1.95))

    # 8. PROGRAM 2 — WHAT IT DOES -------------------------------------------
    s = _blank(prs); _header(s, "Program 2 · residual assembler", "Reconstruct sensitivities offline")
    _bullets(s, Inches(0.55), Inches(1.55), Inches(7.2), Inches(4.4), [
        [[("Links the distributed ", {}), (".obj", {"font": "Consolas", "size": 14}), (" directly — no source, no recompilation of private code.", {})]],
        [[("Replays", {"bold": True, "color": BLUE}), (" the OTI UMAT at each integration point of a converged FE record; for a path-dependent material it marches every increment carrying the state AND its parameter sensitivity.", {})]],
        [[("Assembles", {"bold": True, "color": BLUE}), (" the residual R, its sensitivity ∂R/∂p = Σ Bᵀ (∂σ/∂p) dV, and the tangent K = Σ Bᵀ D B dV.", {})]],
        [[("Solves", {"bold": True, "color": BLUE}), ("  K ∂u/∂p = − ∂R/∂p  for the free DOFs.", {})]],
        [[("Evaluates", {"bold": True, "color": BLUE}), (" each requested ∂q/∂p = ∂q/∂p|explicit + (∂q/∂u) ∂u/∂p.", {})]],
        [[("No new primal solve", {"bold": True, "color": BLUE}), (" in production — the analysis is replayed, not rerun.", {})]],
    ], size=15)
    _codebox(s, Inches(8.0), Inches(1.55), Inches(4.8), Inches(2.2), [
        ("K ∂u/∂p = − ∂R/∂p", ACCENT),
        ("", INK),
        ("∂R/∂p = Σ Bᵀ (∂σ/∂p) dV  − ∂Fext/∂p", INK),
        ("K     = Σ Bᵀ D B dV", INK),
        ("∂q/∂p = ∂q/∂p|exp + (∂q/∂u)·∂u/∂p", GREENT),
    ])
    _takeaway(s, Inches(8.0), Inches(4.1), Inches(4.8),
              "One job file produces the requested sensitivities without rerunning the production analysis.",
              h=Inches(1.3))

    # 9. PROGRAM 2 — THE SENSITIVITY JOB ------------------------------------
    s = _blank(prs); _header(s, "Program 2 · input", "One sensitivity job drives the whole run")
    _bullets(s, Inches(0.55), Inches(1.55), Inches(5.1), Inches(4.3), [
        [[("Seven groups", {"bold": True, "color": BLUE}), (": schema, material, analysis, location, parameters, outputs, result.", {})]],
        [[("material", {"font": "Consolas", "size": 13.5}), (": the OTI .obj + completed contract.", {})]],
        [[("analysis", {"font": "Consolas", "size": 13.5}), (": FE model + converged replay record.", {})]],
        [[("outputs", {"font": "Consolas", "size": 13.5}), (": response, set, component and spatial reduction (average / sum / volume-average).", {})]],
        [[("Request only the sensitivities the collaborator needs.", {"italic": True, "color": MUTED})]],
    ], size=15)
    _codebox(s, Inches(5.9), Inches(1.5), Inches(6.9), Inches(4.7), [
        ('{ "schema": "resasm_sensitivity_job_v1",', INK),
        ('  "material": {"oti_umat":"..._oti.obj",', GREENT),
        ('               "contract":"..._oti.json"},', GREENT),
        ('  "analysis": {"model":"cube.inp",', INK),
        ('               "record":"cube.resrec.h5"},', INK),
        ('  "location": {"step":"LOAD","increment":"last"},', INK),
        ('  "parameters": ["tau0","dG","q","p","gam0","H"],', ACCENT),
        ('  "outputs": [', INK),
        ('    {"type":"integration_point_stress",', INK),
        ('     "region":{"element_set":"DOMAIN"},', INK),
        ('     "component":"S12","reduction":"volume_average"}],', INK),
        ('  "result": {"summary_csv": true} }', INK),
    ])

    # 10. PROGRAM 2 — VERIFICATION METHODOLOGY ------------------------------
    s = _blank(prs); _header(s, "Program 2 · verification", "Checked against full-analysis finite differences")
    _bullets(s, Inches(0.55), Inches(1.55), Inches(12.3), Inches(3.0), [
        [[("Baseline + perturbed FE analyses are run with the ", {}), ("regular UMAT", {"bold": True}), (" — no hard-coded constitutive law anywhere in the reference.", {})]],
        [[("The framework uses the baseline record + OTI .obj to compute R, ∂R/∂p, ∂u/∂p and each requested ∂q/∂p.", {})]],
        [[("Framework sensitivities are compared with ", {}), ("centered finite differences of the perturbed regular-UMAT analyses", {"bold": True}), (".", {})]],
        [[("Passes when global displacement and response sensitivities agree with the independent full-analysis finite differences.", {})]],
    ], size=15)
    _flow(s, Inches(4.7), [("regular UMAT\nbaseline + ±h", False), ("full FE\nanalyses", False),
                           ("central FD\n∂q/∂p", True), ("residual method\n(OTI .obj)", False),
                           ("agree?", True)])
    _takeaway(s, Inches(0.55), Inches(5.8), Inches(12.2),
              "Elastic end-to-end matrix (M1, M2 × displacement/force control): equilibrium ‖R_free‖ ~1e-13, worst sensitivity error < 4e-9 vs full-analysis FD.")

    # 11. PROGRAM 2 — CP WEIGHTED SENSITIVITIES -----------------------------
    s = _blank(prs); _header(s, "Program 2 · results", "Crystal-plasticity sensitivities in one OTI run")
    _figure(s, os.path.join(FIG, "weighted_sigvm_sensitivities.png"), Inches(0.5), Inches(1.55), w=Inches(7.6))
    _bullets(s, Inches(8.35), Inches(1.7), Inches(4.5), Inches(3.8), [
        [[("Single C3D8 element, simple shear, E12: 0 → 0.05 over 50 increments.", {})]],
        [[("All six sensitivities of σ_vM obtained in ", {}), ("one enriched OTI run", {"bold": True, "color": BLUE}), (".", {})]],
        [[("y-axis is a ", {}), ("weighted", {"bold": True}), (" sensitivity — weighted = |p · dσ_vM/dp|, normalized to 100% — so different units compare on one scale.", {})]],
        [[("Glide exponents q, p and activation energy ΔG dominate; hardening H grows with strain; γ₀ stays small.", {})]],
        [[("σ_vM reaches 1.58 GPa at E12 = 0.05.", {"color": MUTED})]],
    ], size=13.5)
    _takeaway(s, Inches(8.35), Inches(5.5), Inches(4.5),
              "Six flow-rule sensitivities, one simulation.", h=Inches(0.8))

    # 12. PROGRAM 2 — OTI vs FD ---------------------------------------------
    s = _blank(prs); _header(s, "Program 2 · results", "OTI matches finite differences — accuracy")
    _figure(s, os.path.join(FIG, "dsigvm_dp_oti_vs_fd.png"), Inches(0.7), Inches(1.55), w=Inches(9.2))
    _bullets(s, Inches(10.1), Inches(1.9), Inches(2.9), Inches(3.5), [
        [[("dσ_vM/dp along the shear path.", {})]],
        [[("OTI (line) vs central FD (points) — indistinguishable.", {})]],
        [[("Worst relative RMSE ", {}), ("2.8e-9", {"bold": True, "color": GREENT}), (" across all six parameters.", {})]],
    ], size=13.5)
    _takeaway(s, Inches(10.1), Inches(5.3), Inches(2.85),
              "OTI is exact to FD truncation — the gold-standard reference.", h=Inches(1.1))

    # 13. PROGRAM 2 — COST --------------------------------------------------
    s = _blank(prs); _header(s, "Program 2 · results", "One enriched run vs many finite-difference runs")
    _figure(s, os.path.join(FIG, "normalized_cost.png"), Inches(0.75), Inches(1.75), h=Inches(4.5))
    _bullets(s, Inches(6.55), Inches(1.9), Inches(6.25), Inches(3.6), [
        [[("Central differences need ", {}), ("2 × 6 + 1 = 13", {"bold": True}), (" full runs for six parameters.", {})]],
        [[("OTI returns all six in ", {}), ("one", {"bold": True, "color": BLUE}), (" enriched march.", {})]],
        [[("The state-alive whole-path march costs ~5.5× a plain UMAT run and returns every sensitivity plus the tangent — ", {}), ("about 2.4× faster than the 13 finite-difference runs", {"bold": True}), (".", {})]],
        [[("An optimized OTI library reduces the per-run overhead substantially (≈1.1× reported elsewhere).", {"color": MUTED})]],
    ], size=14)
    _takeaway(s, Inches(6.55), Inches(5.3), Inches(6.25),
              "Fewer runs, and every derivative is exact — accuracy and cost both favor OTI over finite differences.", h=Inches(1.1))

    # 14. PROGRAM 2 — MESH INDEPENDENCE -------------------------------------
    s = _blank(prs); _header(s, "Program 2 · results", "The residual solve is mesh independent")
    _figure(s, os.path.join(FIG, "weighted_sigvm_sensitivities_4x4x4.png"), Inches(0.5), Inches(1.55), w=Inches(7.6))
    _bullets(s, Inches(8.35), Inches(1.7), Inches(4.5), Inches(3.6), [
        [[("Full residual solve on a ", {}), ("4×4×4 C3D8 mesh", {"bold": True, "color": BLUE}), (" (64 elements, 375 DOF) under controlled simple shear, 50 increments.", {})]],
        [[("Global K and ∂R/∂p assembled; K ∂u/∂p = −∂R/∂p solved for the free DOFs.", {})]],
        [[("Homogeneous deformation ⇒ ‖∂u/∂p‖ = ", {}), ("4.5e-18", {"bold": True, "color": GREENT}), (".", {})]],
        [[("Volume-averaged σ_vM sensitivities match the single element to ", {}), ("1.0e-14", {"bold": True, "color": GREENT}), (".", {})]],
    ], size=13.5)
    _takeaway(s, Inches(8.35), Inches(5.4), Inches(4.5),
              "Sensitivities are mesh independent — a direct check of the global residual assembly.", h=Inches(0.95))

    # 15. SUMMARY -----------------------------------------------------------
    s = _blank(prs); _header(s, "Summary", "Both programs validated, end to end")
    _table(s, Inches(0.55), Inches(1.55), Inches(12.2),
           ["", "What it delivers", "How it is verified", "Result"],
           [[("Program 1", {"bold": True, "color": BLUE}), "UMAT → OTI .obj returning DSIGMA_DP / DSTATEV_DP",
             "vs separately compiled original UMAT + FD", ("18/19 UMATs; sens ≤2.4e-7", {"color": GREENT, "bold": True})],
            [("Program 2", {"bold": True, "color": BLUE}), "Offline residual solve for ∂u/∂p and ∂q/∂p",
             "vs full regular-UMAT finite differences", ("RMSE ~1e-9; mesh-indep.", {"color": GREENT, "bold": True})],
            [("Path dep.", {"bold": True, "color": BLUE}), "DSTATEV_DP chained across increments",
             "multi-increment path vs FD of original", ("M3 ~2e-8 · M5 ~2e-7", {"color": GREENT, "bold": True})]],
           colw=[1.0, 3.6, 3.6, 2.2], fs=12.5, rowh=0.6)
    _bullets(s, Inches(0.55), Inches(4.0), Inches(12.3), Inches(1.9), [
        [[("Material source stays private; the production analysis is unchanged and never rerun in production.", {})]],
        [[("One enriched run returns every requested sensitivity, matching finite differences to machine/FD precision.", {})]],
        [[("Next", {"bold": True, "color": BLUE}), (": drive the actual multi-slip UMAT_PCL through the assembler; add finite-strain and multi-step benchmark cases.", {})]],
    ], size=15)
    _takeaway(s, Inches(0.55), Inches(6.05), Inches(12.2),
              "The framework preserves the original response and produces accurate parameter derivatives — for elasticity, plasticity, viscoplasticity and crystal-plasticity flow.")

    out = os.path.join(RA, "results", "NASA_STRI_residual_assembler.pptx")
    prs.save(out)
    print("wrote", out, "(%d slides)" % len(prs.slides._sldIdLst))
    return out


if __name__ == "__main__":
    build()
